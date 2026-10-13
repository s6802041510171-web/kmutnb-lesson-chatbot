import math
import os
import re
from collections import Counter
from pathlib import Path
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def load_dataset():
    """Load the readable Markdown knowledge base, falling back to Excel."""
    knowledge_path = BASE_DIR / "knowledge_base.md"
    if knowledge_path.exists():
        content = knowledge_path.read_text(encoding="utf-8")
        records = []
        for block in content.split("<!-- END KNOWLEDGE RECORD -->"):
            question_match = re.search(r"\*\*คำถามตัวอย่าง:\*\*\s*(.+)", block)
            keyword_match = re.search(r"\*\*คีย์เวิร์ด:\*\*\s*(.+)", block)
            answer_match = re.search(r"### เนื้อหาจากเอกสาร\s*\n([\s\S]*?)\s*$", block)
            if question_match and answer_match:
                keywords = re.sub(r"[`]", "", keyword_match.group(1)) if keyword_match else ""
                answer = answer_match.group(1).strip()
                if answer:
                    records.append({
                        "question": question_match.group(1).strip(),
                        "answer": answer,
                        "keywords": keywords,
                    })
        if records:
            return records

    """Load the compact Q&A workbook without a heavyweight ML runtime."""
    workbook = load_workbook(BASE_DIR / "dataset.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, ())
    header_names = [str(value or "").strip().lower() for value in headers]

    try:
        question_index = header_names.index("question")
        answer_index = header_names.index("answer")
    except ValueError as exc:
        workbook.close()
        raise RuntimeError("dataset.xlsx must contain Question and Answer columns") from exc

    records = []
    for row in rows:
        question = str(row[question_index] or "").strip() if question_index < len(row) else ""
        # Keep the answer cell verbatim; only the question is normalized for indexing.
        answer = str(row[answer_index]) if answer_index < len(row) and row[answer_index] is not None else ""
        if question and answer.strip():
            records.append({"question": question, "answer": answer, "keywords": ""})

    workbook.close()
    return records


QA_RECORDS = load_dataset()

def clean_text(text: str) -> str:
    text = str(text).lower()
    text = re.sub(r"[^\wก-๙]", "", text)
    return text.strip()


def character_ngrams(text: str):
    """Return compact character features suitable for Thai text without spaces."""
    features = set()
    for size in (2, 3, 4):
        features.update(text[index:index + size] for index in range(len(text) - size + 1))
    return features


ANSWER_RECORDS = []
SEEN_ANSWERS = set()
for record in QA_RECORDS:
    if record["answer"] not in SEEN_ANSWERS:
        ANSWER_RECORDS.append(record)
        SEEN_ANSWERS.add(record["answer"])

STEP_RECORDS = {}
for record in ANSWER_RECORDS:
    match = re.search(r"(?:ขั้นตอนที่|ขันตอนที่)\s*(\d+)", record["answer"], re.IGNORECASE)
    if match:
        STEP_RECORDS.setdefault(int(match.group(1)), record["answer"])

RECORD_FEATURES = [
    character_ngrams(clean_text(item["keywords"] + " " + item["question"] + " " + item["answer"]))
    for item in ANSWER_RECORDS
]
FEATURE_FREQUENCIES = Counter(feature for features in RECORD_FEATURES for feature in features)
FEATURE_IDF = {
    feature: math.log((len(ANSWER_RECORDS) + 1) / (frequency + 1)) + 1
    for feature, frequency in FEATURE_FREQUENCIES.items()
}


def search_qa(query: str):
    # A question asking for the process needs the complete ordered list, not
    # the single passage with the closest wording.
    if re.search(r"ขั้นตอน|กระบวนการ|ลำดับการ", query):
        if STEP_RECORDS:
            parts = ["ขั้นตอนการทำงานตามข้อมูลที่พบในเอกสาร:"]
            missing_steps = [number for number in range(1, max(STEP_RECORDS) + 1) if number not in STEP_RECORDS]
            if missing_steps:
                missing = ", ".join(str(number) for number in missing_steps)
                parts.append(f"หมายเหตุ: ไม่พบเนื้อหาขั้นตอนที่ {missing} ในฐานความรู้ จึงไม่เติมข้อมูลขึ้นเอง")
            for number, answer in sorted(STEP_RECORDS.items()):
                match = re.search(
                    rf"(?:ขั้นตอนที่|ขันตอนที่)\s*{number}\s*([^\n]*)",
                    answer,
                    re.IGNORECASE,
                )
                title = match.group(1).strip(" .:-") if match else ""
                body = re.sub(r"\s+", " ", answer).strip()
                keyword = re.sub(r"[^\wก-๙+#.-]", " ", title).strip()
                parts.append(f"\nขั้นตอนที่ {number}: {title}\nคีย์เวิร์ด: {keyword}\nรายละเอียดจากเอกสาร: {body}")
            return "\n".join(parts)

    q_clean = clean_text(query)
    query_features = character_ngrams(q_clean)
    if not query_features:
        return None

    ranked = []
    query_weight = sum(FEATURE_IDF.get(feature, 1) for feature in query_features)
    for record, features in zip(ANSWER_RECORDS, RECORD_FEATURES):
        # Match against both the indexed question and its source passage. The
        # question often carries page/section context absent from the passage.
        score = sum(FEATURE_IDF.get(feature, 1) for feature in query_features & features) / query_weight
        if score:
            ranked.append((score, record))

    ranked.sort(key=lambda item: item[0], reverse=True)
    # Short Thai questions and section titles have little character overlap
    # with long passages, so use a softer threshold while still requiring
    # meaningful evidence from the document index.
    if ranked and ranked[0][0] >= 0.16:
        return ranked[0][1]["answer"]
    return None

class ChatRequest(BaseModel):
    message: str

@app.get("/")
def read_root():
    return FileResponse(BASE_DIR / "index.html")

@app.get("/style.css", include_in_schema=False)
def read_stylesheet():
    return FileResponse(BASE_DIR / "style.css", media_type="text/css")

@app.get("/health")
def health_check():
    return {
        "status": "ok",
        "total_records_in_dataset": len(QA_RECORDS),
        "source": "Markdown knowledge base" if (BASE_DIR / "knowledge_base.md").exists() else "Excel dataset"
    }

@app.post("/chat")
async def chat_endpoint(req: ChatRequest):
    user_msg = req.message.strip()
    answer = search_qa(user_msg)
    if answer is not None:
        return {"source": "dataset", "reply": answer}
    return {"source": "not_found", "reply": "ไม่พบคำตอบที่ตรงกับข้อมูลในเอกสาร DatasetภาษาC"}

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 10000))
    uvicorn.run("app:app", host="0.0.0.0", port=port)
