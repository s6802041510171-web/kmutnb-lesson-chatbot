"""Build a readable Markdown knowledge base from dataset.xlsx."""

import re
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
END_MARKER = "<!-- END KNOWLEDGE RECORD -->"


def extract_title(answer: str) -> str:
    step = re.search(r"(?:ขั้นตอนที่|ขันตอนที่)\s*\d+[^\n]*", answer, re.I)
    if step:
        return step.group(0).strip()[:140]
    for line in answer.splitlines():
        line = line.strip()
        if re.match(r"\d+(?:\.\d+)+\.?\s+", line):
            return line[:140]
    for line in answer.splitlines():
        line = line.strip()
        if line:
            return line[:90]
    return "เนื้อหาจากเอกสาร"


def extract_keywords(question: str, answer: str) -> list[str]:
    title = extract_title(answer)
    values = []
    section = re.search(r"\b\d+(?:\.\d+)+\b", answer)
    if section:
        values.append(section.group(0))
    step = re.search(r"(?:ขั้นตอนที่|ขันตอนที่)\s*\d+\s*([^\n]*)", answer, re.I)
    if step:
        title = step.group(1)
    elif section:
        title = re.sub(r"^\s*\d+(?:\.\d+)+\.?\s*", "", title)
    title = re.sub(r"[\s,;:()]+", " ", title).strip()
    if title:
        values.append(title[:100])
    stopwords = {"เอกสาร", "page", "หน้า", "ช่วง", "dataset", "question", "chat", "bot", "language"}
    for token in re.findall(r"[A-Za-z][A-Za-z0-9+#.-]*", title + " " + question):
        if token.lower() in stopwords:
            continue
        if len(token) > 1 and token.lower() not in {v.lower() for v in values}:
            values.append(token)
        if len(values) >= 10:
            break
    return list(dict.fromkeys(values))


def main():
    workbook = load_workbook(BASE_DIR / "dataset.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().lower() for value in next(rows, ())]
    question_index = headers.index("question")
    answer_index = headers.index("answer")

    records = []
    seen = set()
    for row in rows:
        question = str(row[question_index] or "").strip()
        answer = str(row[answer_index] or "").strip()
        if question and answer and answer not in seen:
            seen.add(answer)
            records.append((question, answer))
    workbook.close()

    output = [
        "# ฐานความรู้: บทเรียนภาษา C",
        "",
        "ไฟล์นี้สร้างจากชีต Question/Answer ใน `dataset.xlsx` ใช้หัวข้อและคีย์เวิร์ดช่วยค้นหา",
        "เนื้อหาคำตอบคงข้อความจากเอกสารต้นทาง",
        "",
    ]
    for number, (question, answer) in enumerate(records, 1):
        keywords = extract_keywords(question, answer)
        output.extend([
            f"## {extract_title(answer)}",
            f"**คีย์เวิร์ด:** {', '.join(f'`{item}`' for item in keywords)}",
            f"**คำถามตัวอย่าง:** {question}",
            "",
            "### เนื้อหาจากเอกสาร",
            answer,
            "",
            END_MARKER,
            "",
        ])
    (BASE_DIR / "knowledge_base.md").write_text("\n".join(output), encoding="utf-8")
    print(f"สร้าง knowledge_base.md จำนวน {len(records)} รายการ")


if __name__ == "__main__":
    main()
